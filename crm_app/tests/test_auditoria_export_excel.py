"""Testes da exportação Excel da fila de auditoria."""
from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from crm_app.models import Cliente, StatusCRM, Venda
from crm_app.services.auditoria_export_service import HEADERS
from usuarios.models import Perfil, Usuario

URL_EXPORT = '/api/crm/vendas/exportar-auditoria-excel/'


class ExportarAuditoriaExcelTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.perfil_aud = Perfil.objects.create(cod_perfil='AUD', nome='Auditoria')
        cls.perfil_vend = Perfil.objects.create(cod_perfil='VEND', nome='Vendedor')
        cls.auditor = Usuario.objects.create_user(
            username='aud_export',
            password='SenhaSegura123',
            first_name='Ana',
            last_name='Auditora',
            perfil=cls.perfil_aud,
        )
        cls.vendedor = Usuario.objects.create_user(
            username='vend_export',
            password='SenhaSegura123',
            first_name='Carlos',
            last_name='Vendedor',
            perfil=cls.perfil_vend,
        )
        cls.cliente = Cliente.objects.create(
            cpf_cnpj='12345678901',
            nome_razao_social='CLIENTE EXPORT AUDITORIA',
        )
        cls.st_aberto = StatusCRM.objects.create(
            nome='SEM TRATAMENTO', tipo='Tratamento', estado='ABERTO',
        )
        cls.st_pend = StatusCRM.objects.create(
            nome='PENDENTE VENDEDOR', tipo='Tratamento', estado='ABERTO',
        )
        cls.st_fechado = StatusCRM.objects.create(
            nome='CADASTRADA', tipo='Tratamento', estado='FECHADO',
        )
        cls.st_esteira = StatusCRM.objects.create(
            nome='AGENDADO', tipo='Esteira', estado='ABERTO',
        )

    def _criar_venda(self, **kwargs) -> Venda:
        dados = {
            'vendedor': self.vendedor,
            'cliente': self.cliente,
            'status_tratamento': self.st_aberto,
            'ativo': True,
            'telefone1': '21988887777',
            'cidade': 'Rio de Janeiro',
            'estado': 'RJ',
        }
        dados.update(kwargs)
        return Venda.objects.create(**dados)

    def _ids_planilha(self, content: bytes) -> list[int]:
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        ids = []
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0] is not None:
                ids.append(int(row[0]))
        return ids

    def test_exige_autenticacao(self):
        r = self.client.get(URL_EXPORT)
        self.assertEqual(r.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_vendedor_recebe_403(self):
        self.client.force_authenticate(user=self.vendedor)
        r = self.client.get(URL_EXPORT)
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    def test_exporta_todos_os_pendentes_sem_paginacao(self):
        v1 = self._criar_venda()
        v2 = self._criar_venda(telefone1='21977776666')
        fechada = self._criar_venda(status_tratamento=self.st_fechado)
        na_esteira = self._criar_venda(status_esteira=self.st_esteira)

        self.client.force_authenticate(user=self.auditor)
        r = self.client.get(URL_EXPORT)
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            r['Content-Type'],
        )
        self.assertIn('attachment; filename="auditoria_pendentes_', r['Content-Disposition'])

        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        self.assertEqual([c.value for c in ws[1]], HEADERS)
        ids = self._ids_planilha(r.content)
        self.assertIn(v1.id, ids)
        self.assertIn(v2.id, ids)
        self.assertNotIn(fechada.id, ids)
        self.assertNotIn(na_esteira.id, ids)

        linha_v1 = next(
            row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] == v1.id
        )
        self.assertEqual(linha_v1[3], 'CLIENTE EXPORT AUDITORIA')
        self.assertEqual(linha_v1[4], '12345678901')
        self.assertEqual(linha_v1[5], 'Carlos Vendedor')
        self.assertEqual(linha_v1[7], 'SEM TRATAMENTO')
        self.assertEqual(linha_v1[11], 'Pendente')
        self.assertEqual(linha_v1[12], '21988887777')
        self.assertEqual(linha_v1[15], 'Rio de Janeiro')
        self.assertEqual(linha_v1[16], 'RJ')

    def test_filtra_por_status_tratamento(self):
        aberta = self._criar_venda()
        self._criar_venda(status_tratamento=self.st_pend)

        self.client.force_authenticate(user=self.auditor)
        r = self.client.get(URL_EXPORT, {'status_tratamento_id': self.st_aberto.id})
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(self._ids_planilha(r.content), [aberta.id])

    def test_filtra_por_periodo(self):
        hoje = timezone.localdate()
        recente = self._criar_venda()
        antiga = self._criar_venda()
        Venda.objects.filter(pk=antiga.pk).update(
            data_criacao=timezone.now() - timedelta(days=10),
        )

        self.client.force_authenticate(user=self.auditor)
        r = self.client.get(URL_EXPORT, {
            'data_inicio': hoje.isoformat(),
            'data_fim': hoje.isoformat(),
        })
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(self._ids_planilha(r.content), [recente.id])
