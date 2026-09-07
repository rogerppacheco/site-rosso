"""Exportação Excel da fila de auditoria (vendas pendentes)."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

import openpyxl
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from crm_app.models import Venda

HEADERS: list[str] = [
    'ID',
    'Pedido PAP',
    'Data Venda',
    'Cliente',
    'CPF/CNPJ',
    'Vendedor',
    'Produto',
    'Status Tratamento',
    'Última Edição por',
    'Data Última Edição',
    'Auditor',
    'Biometria',
    'Telefone 1',
    'Telefone 2',
    'O.S.',
    'Cidade',
    'UF',
    'Data Agendamento',
    'Turno',
    'Observações',
]


def _fmt_dt(valor: datetime | None) -> str:
    if not valor:
        return ''
    if timezone.is_aware(valor):
        valor = timezone.localtime(valor)
    return valor.strftime('%d/%m/%Y %H:%M')


def _fmt_date(valor: date | None) -> str:
    if not valor:
        return ''
    return valor.strftime('%d/%m/%Y')


def _rotulo_biometria(aprovada: bool | None) -> str:
    if aprovada is True:
        return 'Aprovada'
    if aprovada is False:
        return 'Não aprovada'
    return 'Pendente'


def _nome_usuario(user: Any) -> str:
    if not user:
        return ''
    nome = (user.get_full_name() or '').strip()
    return nome or user.username or ''


def _linha_venda(venda: Venda) -> list[Any]:
    cliente = venda.cliente
    vendedor = venda.vendedor
    return [
        venda.id,
        venda.pedido_pap or '',
        _fmt_dt(venda.data_criacao),
        cliente.nome_razao_social if cliente else '',
        cliente.cpf_cnpj if cliente else '',
        _nome_usuario(vendedor) or (venda.vendedor_matricula_pap or ''),
        venda.plano.nome if venda.plano else '',
        venda.status_tratamento.nome if venda.status_tratamento else '',
        _nome_usuario(venda.editado_por),
        _fmt_dt(venda.data_ultima_alteracao),
        _nome_usuario(venda.auditor_atual),
        _rotulo_biometria(venda.biometria_aprovada),
        venda.telefone1 or '',
        venda.telefone2 or '',
        venda.ordem_servico or '',
        venda.cidade or '',
        venda.estado or '',
        _fmt_date(venda.data_agendamento),
        venda.get_periodo_agendamento_display() or '',
        venda.observacoes or '',
    ]


def montar_xlsx_pendentes_auditoria(vendas: Iterable[Venda]) -> tuple[bytes, str]:
    """Gera planilha com todas as vendas da fila de auditoria (sem paginação)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Auditoria'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='375A7F')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.freeze_panes = 'A2'

    for venda in vendas:
        ws.append(_linha_venda(venda))

    larguras = [10, 16, 18, 32, 16, 18, 22, 22, 18, 18, 18, 14, 16, 16, 16, 18, 8, 16, 12, 40]
    for idx, width in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ultima = ws.max_row or 1
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}{ultima}'

    buf = BytesIO()
    wb.save(buf)
    nome = f"auditoria_pendentes_{timezone.localtime().strftime('%Y%m%d_%H%M')}.xlsx"
    return buf.getvalue(), nome
